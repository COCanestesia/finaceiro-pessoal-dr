from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font,Alignment
from .models import ReportTable

def export_xlsx(table:ReportTable,destination:Path)->Path:
    destination=Path(destination);destination.parent.mkdir(parents=True,exist_ok=True);wb=Workbook();ws=wb.active;ws.title='Relatório';ws.append([table.title]);ws['A1'].font=Font(bold=True,size=16);ws.append([table.filters_text]);ws.append(list(table.columns))
    for cell in ws[3]:cell.font=Font(bold=True)
    for row in table.rows:ws.append(list(row))
    if table.total_cents is not None:ws.append([]);ws.append(['Total',table.total_cents/100]);ws.cell(ws.max_row,2).number_format='R$ #,##0.00'
    for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=min(45,max(12,max(len(str(c.value or '')) for c in col)+2))
    wb.save(destination);return destination
