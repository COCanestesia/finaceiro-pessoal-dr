from __future__ import annotations

import calendar
import sqlite3
from pathlib import Path


def monthly_summary(connection: sqlite3.Connection, year: int, month: int) -> dict:
    """Financial entries by competence, excluding canceled and deleted entries."""
    if not (1 <= month <= 12 and 2000 <= year <= 2200):
        raise ValueError("Período inválido.")
    prefix = f"{year:04d}-{month:02d}-%"
    rows = connection.execute(
        "SELECT entry_type, status, COUNT(*) AS quantity, COALESCE(SUM(amount_cents),0) AS amount "
        "FROM financial_entry WHERE competence_date LIKE ? AND deleted_at IS NULL "
        "AND status <> 'CANCELADO' GROUP BY entry_type,status ORDER BY entry_type,status",
        (prefix,),
    ).fetchall()
    totals = {"RECEITA": 0, "DESPESA": 0}
    details = []
    for row in rows:
        value = int(row["amount"])
        details.append({"type": row["entry_type"], "status": row["status"],
                        "quantity": row["quantity"], "amount_cents": value})
        if row["entry_type"] in totals:
            totals[row["entry_type"]] += value
    return {"year": year, "month": month, "income_cents": totals["RECEITA"],
            "expense_cents": totals["DESPESA"],
            "result_cents": totals["RECEITA"] - totals["DESPESA"], "details": details}


def export_excel(connection: sqlite3.Connection, path: Path, year: int, month: int) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    summary = monthly_summary(connection, year, month)
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumo"
    ws.append(["Financeiro Pessoal DR", f"{month:02d}/{year}"])
    for name, key in (("Receitas", "income_cents"), ("Despesas", "expense_cents"),
                      ("Resultado", "result_cents")):
        ws.append([name, summary[key] / 100])
        ws.cell(ws.max_row, 2).number_format = '"R$" #,##0.00;[Red]("R$" #,##0.00)'
    detail = wb.create_sheet("Detalhamento")
    detail.append(["Tipo", "Status", "Quantidade", "Valor (R$)"])
    for item in summary["details"]:
        detail.append([item["type"], item["status"], item["quantity"], item["amount_cents"] / 100])
        detail.cell(detail.max_row,4).number_format = '"R$" #,##0.00'
    for sheet in (ws,detail):
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="20354B")
        for col in "ABCD":
            sheet.column_dimensions[col].width = 24
    dest = Path(path)
    if dest.exists():
        raise FileExistsError("Escolha outro nome para o relatório.")
    wb.save(dest)
    return dest


def export_pdf(connection: sqlite3.Connection, path: Path, year: int, month: int) -> Path:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen.canvas import Canvas
    summary = monthly_summary(connection,year,month)
    dest = Path(path)
    if dest.exists():
        raise FileExistsError("Escolha outro nome para o relatório.")
    canvas = Canvas(str(dest), pagesize=A4)
    canvas.setTitle("Resumo financeiro pessoal")
    canvas.setFont("Helvetica-Bold", 18)
    canvas.drawString(48, 790, "Financeiro Pessoal DR")
    canvas.setFont("Helvetica", 12)
    canvas.drawString(48, 765, f"Resumo mensal - {month:02d}/{year}")
    def currency(cents):
        return f"R$ {cents / 100:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    for index, (label,key) in enumerate((("Receitas","income_cents"),("Despesas","expense_cents"),("Resultado","result_cents"))):
        canvas.drawString(48, 725-index*28, f"{label}: {currency(summary[key])}")
    canvas.setFont("Helvetica-Bold", 10)
    canvas.drawString(48, 610, "Tipo")
    canvas.drawString(160, 610, "Status")
    canvas.drawString(300, 610, "Qtd.")
    canvas.drawString(365, 610, "Valor")
    canvas.setFont("Helvetica", 10)
    y = 590
    for item in summary["details"]:
        if y < 70:
            canvas.showPage()
            y = 780
        for x, value in ((48,item["type"]),(160,item["status"]),(300,str(item["quantity"])),(365,currency(item["amount_cents"]))):
            canvas.drawString(x,y,value)
        y -= 18
    canvas.save()
    return dest
