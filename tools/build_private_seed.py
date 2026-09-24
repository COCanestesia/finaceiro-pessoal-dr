from __future__ import annotations

import argparse
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
import hashlib
import json
from pathlib import Path
import re

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel


REQUIRED_HEADERS = (
    "TITULAR",
    "DATA",
    "DESCRIÇÃO",
    "CONTA",
    "VALOR",
    "CATEGORIA",
    "TIPO DE DESPESA",
    "CLASSIFICAÇÃO",
)

_CATEGORY_ALIASES = {
    "CARTAO DE CREDITO": "CARTÃO DE CRÉDITO",
}

_TYPE_ALIASES = {
    "DESPESA DIARIA": "DESPESA DIÁRIA",
}


def _clean_text(value) -> str | None:
    if value is None:
        return None
    text = re.sub(r"\s+", " ", str(value).strip())
    return text or None


def _canonical_category(value) -> str | None:
    cleaned = _clean_text(value)
    if cleaned is None:
        return None
    return _CATEGORY_ALIASES.get(cleaned.upper(), cleaned)


def _canonical_raw_type(value) -> str | None:
    cleaned = _clean_text(value)
    if cleaned is None:
        return None
    key = cleaned.upper()
    return _TYPE_ALIASES.get(key, key)


def _classification(value) -> str:
    cleaned = _clean_text(value)
    if cleaned is None:
        raise ValueError("CLASSIFICAÇÃO vazia na planilha.")
    result = cleaned.upper()
    if result not in {"RECEITA", "DESPESA"}:
        raise ValueError(f"CLASSIFICAÇÃO desconhecida: {cleaned}")
    return result


def _iso_date(value, *, epoch) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        converted = from_excel(value, epoch=epoch)
        if isinstance(converted, datetime):
            converted = converted.date()
        if isinstance(converted, date):
            return converted.isoformat()
    if isinstance(value, str):
        text = value.strip()
        try:
            return date.fromisoformat(text[:10]).isoformat()
        except ValueError:
            pass
    raise ValueError(f"DATA inválida na planilha: {value!r}")


def _amount_to_cents(value) -> tuple[int, str | None]:
    if value is None or (isinstance(value, str) and not value.strip()):
        return 0, "MISSING_AMOUNT"
    try:
        amount = Decimal(str(value)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    except (InvalidOperation, ValueError) as exc:
        raise ValueError(f"VALOR inválido na planilha: {value!r}") from exc
    cents = int(amount * 100)
    if cents < 0:
        raise ValueError("VALOR negativo não é aceito na carga inicial.")
    return cents, None


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def build_seed(
    input_path: Path,
    output_path: Path,
    *,
    seed_id: str,
    sheet_name: str = "Entrada_dados",
) -> dict:
    input_path = Path(input_path)
    output_path = Path(output_path)
    seed_id = (_clean_text(seed_id) or "")
    if not seed_id:
        raise ValueError("seed_id é obrigatório.")

    workbook = load_workbook(input_path, data_only=True, read_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"A aba {sheet_name!r} não existe na planilha.")
        sheet = workbook[sheet_name]
        iterator = sheet.iter_rows(values_only=True)
        try:
            header_values = next(iterator)
        except StopIteration as exc:
            raise ValueError("A planilha está vazia.") from exc

        header_map: dict[str, int] = {}
        for index, value in enumerate(header_values):
            cleaned = _clean_text(value)
            if cleaned:
                header_map[cleaned.upper()] = index
        missing = [header for header in REQUIRED_HEADERS if header not in header_map]
        if missing:
            raise ValueError("Colunas obrigatórias ausentes: " + ", ".join(missing))

        records: list[dict] = []
        income_count = 0
        expense_count = 0
        income_cents = 0
        expense_cents = 0
        for source_row, row in enumerate(iterator, start=2):
            def value(header: str):
                index = header_map[header]
                return row[index] if index < len(row) else None

            significant = [
                value("TITULAR"),
                value("DATA"),
                value("DESCRIÇÃO"),
                value("VALOR"),
                value("CLASSIFICAÇÃO"),
            ]
            if all(item is None or (isinstance(item, str) and not item.strip()) for item in significant):
                continue

            holder = _clean_text(value("TITULAR"))
            description = _clean_text(value("DESCRIÇÃO"))
            if not holder:
                raise ValueError(f"TITULAR vazio na linha {source_row}.")
            if not description:
                raise ValueError(f"DESCRIÇÃO vazia na linha {source_row}.")

            classification = _classification(value("CLASSIFICAÇÃO"))
            amount_cents, warning = _amount_to_cents(value("VALOR"))
            record = {
                "source_row": source_row,
                "holder": holder,
                "date": _iso_date(value("DATA"), epoch=workbook.epoch),
                "description": description,
                "payment_method": _clean_text(value("CONTA")),
                "amount_cents": amount_cents,
                "category": _canonical_category(value("CATEGORIA")),
                "raw_type": _canonical_raw_type(value("TIPO DE DESPESA")),
                "classification": classification,
                "warning": warning,
            }
            records.append(record)
            if classification == "RECEITA":
                income_count += 1
                income_cents += amount_cents
            else:
                expense_count += 1
                expense_cents += amount_cents
    finally:
        workbook.close()

    payload = {
        "schema_version": 1,
        "seed_id": seed_id,
        "source_label": input_path.name,
        "source_sha256": _sha256(input_path),
        "expected": {
            "record_count": len(records),
            "income_count": income_count,
            "expense_count": expense_count,
            "income_cents": income_cents,
            "expense_cents": expense_cents,
        },
        "records": records,
    }
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return payload


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Gera um seed privado para o Financeiro Pessoal do Dr.")
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--seed-id", required=True)
    parser.add_argument("--sheet", default="Entrada_dados")
    args = parser.parse_args(argv)
    build_seed(args.input, args.output, seed_id=args.seed_id, sheet_name=args.sheet)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
